import os
import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from src.chunking import DocumentChunk, chunk_documents, chunk_text
from src.document_loader import load_docx, load_pdf, load_pptx, load_txt
from src.generator import AnswerGenerator
from src.hybrid_search import HybridSearcher
from src.keyword_search import KeywordSearcher
from src.reranker import CrossEncoderReranker


def make_minimal_text_pdf(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)

    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font_reference}
            )
        }
    )

    escaped_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = DecodedStreamObject()
    content.set_data(
        f"BT /F1 12 Tf 72 720 Td ({escaped_text}) Tj ET".encode("latin-1")
    )
    page[NameObject("/Contents")] = writer._add_object(content)

    output = BytesIO()
    writer.write(output)
    return output.getvalue()


class DocumentExtractionParityTests(unittest.TestCase):
    def test_txt_extraction_uses_utf8(self):
        self.assertEqual(load_txt("RAG parity ✓".encode("utf-8")), "RAG parity ✓")

    def test_docx_extraction_preserves_nonempty_paragraph_order(self):
        document = Document()
        document.add_paragraph("First paragraph")
        document.add_paragraph("")
        document.add_paragraph("Second paragraph")
        output = BytesIO()
        document.save(output)

        self.assertEqual(
            load_docx(output.getvalue()),
            "First paragraph\nSecond paragraph",
        )

    def test_pptx_extraction_preserves_slide_title_and_content(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Parity Deck"
        slide.placeholders[1].text = "First point\nSecond point"
        output = BytesIO()
        presentation.save(output)

        self.assertEqual(
            load_pptx(output.getvalue()),
            (
                "[Slide 1]\n"
                "Title: Parity Deck\n"
                "Content:\n"
                "- First point\n"
                "- Second point"
            ),
        )

    def test_pdf_extraction_reads_a_minimal_generated_text_fixture(self):
        extracted = load_pdf(make_minimal_text_pdf("Minimal PDF parity fixture"))

        self.assertEqual(extracted.strip(), "Minimal PDF parity fixture")


class ChunkingParityTests(unittest.TestCase):
    def test_default_chunks_use_120_words_with_30_word_overlap(self):
        words = [f"word{index}" for index in range(210)]

        chunks = chunk_text(" ".join(words), "notes.txt", start_chunk_id=1)

        self.assertEqual(len(chunks), 3)
        self.assertEqual(len(chunks[0].text.split()), 120)
        self.assertEqual(chunks[0].text.split()[-30:], chunks[1].text.split()[:30])
        self.assertEqual(chunks[1].text.split()[0], "word90")
        self.assertEqual(chunks[2].text.split()[0], "word180")

    def test_chunk_ids_are_sequential_and_filenames_are_preserved(self):
        documents = [
            {"source": "first.txt", "text": " ".join(["first"] * 130)},
            {"source": "second.txt", "text": " ".join(["second"] * 10)},
        ]

        chunks = chunk_documents(documents)

        self.assertEqual([chunk.chunk_id for chunk in chunks], [1, 2, 3])
        self.assertEqual(
            [chunk.source for chunk in chunks],
            ["first.txt", "first.txt", "second.txt"],
        )


class KeywordRetrievalParityTests(unittest.TestCase):
    def test_bm25_tokenization_is_lowercase_regex_word_tokenization(self):
        searcher = KeywordSearcher([])

        self.assertEqual(
            searcher._tokenize("RAG, BM25! Version_2"),
            ["rag", "bm25", "version_2"],
        )

    def test_bm25_retrieval_ranks_the_lexical_match_first(self):
        chunks = [
            DocumentChunk(1, "fruit.txt", "apple orchard fruit"),
            DocumentChunk(2, "code.txt", "python code function"),
            DocumentChunk(3, "recipe.txt", "apple pie recipe"),
            DocumentChunk(4, "data.txt", "database storage query"),
        ]

        results = KeywordSearcher(chunks).search("python function", top_k=2)

        self.assertEqual(len(results), 2)
        self.assertEqual(results[0]["chunk_id"], 2)
        self.assertEqual(results[0]["source"], "code.txt")
        self.assertGreater(results[0]["score"], results[1]["score"])


class HybridRetrievalParityTests(unittest.TestCase):
    def test_default_weighted_combination_and_ranking_are_stable(self):
        chunks = [
            DocumentChunk(1, "one.txt", "one"),
            DocumentChunk(2, "two.txt", "two"),
            DocumentChunk(3, "three.txt", "three"),
        ]

        class StubSemanticSearcher:
            def __init__(self, supplied_chunks):
                self.chunks = supplied_chunks

            def search(self, query, top_k=5):
                scores = {1: 0.9, 2: 0.5, 3: 0.1}
                return [
                    {
                        "chunk_id": chunk.chunk_id,
                        "source": chunk.source,
                        "text": chunk.text,
                        "score": scores[chunk.chunk_id],
                    }
                    for chunk in self.chunks[:top_k]
                ]

        class StubKeywordSearcher:
            def __init__(self, supplied_chunks):
                self.chunks = supplied_chunks

            def search(self, query, top_k=5):
                scores = {1: 0.0, 2: 10.0, 3: 0.0}
                return [
                    {
                        "chunk_id": chunk.chunk_id,
                        "source": chunk.source,
                        "text": chunk.text,
                        "score": scores[chunk.chunk_id],
                    }
                    for chunk in self.chunks[:top_k]
                ]

        with (
            patch("src.hybrid_search.SemanticSearcher", StubSemanticSearcher),
            patch("src.hybrid_search.KeywordSearcher", StubKeywordSearcher),
        ):
            searcher = HybridSearcher(chunks)
            first_run = searcher.search("query", top_k=3)
            second_run = searcher.search("query", top_k=3)

        self.assertEqual([result["chunk_id"] for result in first_run], [2, 1, 3])
        self.assertEqual(
            [result["chunk_id"] for result in first_run],
            [result["chunk_id"] for result in second_run],
        )
        self.assertAlmostEqual(first_run[0]["semantic_score"], 0.5)
        self.assertAlmostEqual(first_run[0]["keyword_score"], 1.0)
        self.assertAlmostEqual(first_run[0]["hybrid_score"], 0.675)
        self.assertAlmostEqual(first_run[1]["hybrid_score"], 0.65)

    def test_equal_scores_currently_normalize_to_one(self):
        """Parity lock: equal scores, including all-zero scores, currently become 1.0."""
        searcher = HybridSearcher.__new__(HybridSearcher)
        equal_score_results = [
            {"chunk_id": 1, "score": 0.0},
            {"chunk_id": 2, "score": 0.0},
        ]

        normalized = searcher._normalize_scores(equal_score_results)

        self.assertEqual(normalized, {1: 1.0, 2: 1.0})


class RerankingParityTests(unittest.TestCase):
    def test_reranker_preserves_result_fields_sorts_scores_and_applies_top_k(self):
        class StubCrossEncoder:
            def __init__(self, model_name):
                self.model_name = model_name

            def predict(self, pairs):
                self.pairs = pairs
                return [0.25, 0.9, -0.1]

        input_results = [
            {"chunk_id": 1, "source": "one.txt", "text": "first", "hybrid_score": 0.8},
            {"chunk_id": 2, "source": "two.txt", "text": "second", "hybrid_score": 0.7},
            {"chunk_id": 3, "source": "three.txt", "text": "third", "hybrid_score": 0.6},
        ]

        with patch("src.reranker.CrossEncoder", StubCrossEncoder):
            reranker = CrossEncoderReranker()
            output_results = reranker.rerank("question", input_results, top_k=2)

        self.assertEqual([result["chunk_id"] for result in output_results], [2, 1])
        self.assertEqual(output_results[0]["source"], "two.txt")
        self.assertEqual(output_results[0]["text"], "second")
        self.assertEqual(output_results[0]["hybrid_score"], 0.7)
        self.assertEqual(output_results[0]["rerank_score"], 0.9)
        self.assertNotIn("rerank_score", input_results[0])


class AnswerGenerationParityTests(unittest.TestCase):
    def setUp(self):
        self.results = [
            {
                "chunk_id": 1,
                "source": "guide.txt",
                "text": (
                    "BM25 keyword search retrieves lexically matching passages. "
                    "Reranking improves their order."
                ),
            }
        ]

    def test_retrieval_only_answer_is_deterministic(self):
        generator = AnswerGenerator(provider="none")

        first_answer = generator.generate_answer(
            "How does BM25 keyword search retrieve passages?",
            self.results,
        )
        second_answer = generator.generate_answer(
            "How does BM25 keyword search retrieve passages?",
            self.results,
        )

        expected = (
            "BM25 keyword search retrieves lexically matching passages. "
            "Reranking improves their order.\n\n"
            "Sources used:\n"
            "1. Topic: BM25 keyword search — File: guide.txt"
        )
        self.assertEqual(first_answer, expected)
        self.assertEqual(second_answer, expected)

    def test_source_construction_uses_topic_heuristics_and_filename(self):
        generator = AnswerGenerator(provider="none")
        source_results = [
            {"source": "resume.docx", "text": "Education and coursework"},
            {"source": "rag.pdf", "text": "Embedding vectors support semantic search"},
        ]

        self.assertEqual(
            generator._source_lines(source_results),
            [
                "1. Topic: Education — File: resume.docx",
                "2. Topic: Embeddings and semantic search — File: rag.pdf",
            ],
        )

    def test_summary_like_question_routes_to_full_document_summary(self):
        generator = AnswerGenerator(provider="none")

        with patch.object(
            generator,
            "_generate_full_document_summary",
            return_value="SUMMARY",
        ) as summary_generator:
            answer = generator.generate_answer("Please summarize this document", self.results)

        self.assertEqual(answer, "SUMMARY")
        summary_generator.assert_called_once_with(
            "Please summarize this document",
            self.results,
        )

    def _assert_missing_key_falls_back(self, provider):
        with (
            patch("src.generator.load_dotenv"),
            patch.dict(os.environ, {}, clear=True),
        ):
            generator = AnswerGenerator(provider=provider)

        with patch.object(
            generator,
            "_generate_provider_answer",
            side_effect=AssertionError("Provider API must not be called without a key"),
        ):
            answer = generator.generate_answer(
                "How does BM25 keyword search retrieve passages?",
                self.results,
            )

        self.assertTrue(answer.startswith(AnswerGenerator.MISSING_PROVIDER_KEY_NOTE))
        self.assertIn("BM25 keyword search retrieves", answer)
        self.assertIn("Sources used:", answer)

    def test_openai_without_api_key_uses_retrieval_fallback(self):
        self._assert_missing_key_falls_back("openai")

    def test_anthropic_without_api_key_uses_retrieval_fallback(self):
        self._assert_missing_key_falls_back("anthropic")

    def test_gemini_without_api_key_uses_retrieval_fallback(self):
        self._assert_missing_key_falls_back("gemini")


if __name__ == "__main__":
    unittest.main()
